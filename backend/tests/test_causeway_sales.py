import io
from dataclasses import asdict
from datetime import datetime

import msoffcrypto.format.ooxml
import openpyxl
import pytest

from app.workflows.base import InputFile, RunContext, WorkflowInputError
from app.workflows.finance.causeway_sales.logic import ZOHO_COLUMNS
from app.workflows.finance.causeway_sales.workflow import CausewaySalesReconciliation

AFS_PASSWORD = "4651"


def make_afs(
    day: str = "18-OCT-2026",
    password: str = AFS_PASSWORD,
    total: float = 91,
    later: tuple[str, ...] = (),
) -> bytes:
    """`later` moves the last sales to these dates, one each (e.g. next-day transactions)."""
    """Same layout as the AFS settlement report: 10 sales totalling 91.000 on one terminal."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MerchantAUBTerminalWise"
    ws.append(["Terminal ID:03221330"])
    ws.append(
        [
            "Tran Date",
            None,
            "Batch",
            "RRN",
            "Card No",
            "Tran Type",
            "Auth Code",
            "Track  Id",
            "Tran Amt",
            "Commission",
            "VAT",
            "Net",
        ]
    )
    sales = [("8.500", "-0.170", "-0.017", "8.313")] * 7 + [
        ("9.500", "-0.190", "-0.019", "9.291"),
        ("11.000", "-0.220", "-0.022", "10.758"),
        ("11.000", "-0.220", "-0.022", "10.758"),
    ]
    days = [day] * (len(sales) - len(later)) + list(later)
    for (amt, com, vat, net), tran_day in zip(sales, days, strict=True):
        ws.append(
            [
                tran_day,
                None,
                "999",
                "629113075038",
                "440227******1751",
                "Purchase",
                "580740",
                "",
                f"{amt} BHD",
                com,
                vat,
                net,
            ]
        )
    ws.append([None, "Total", None, None, None, None, None, None, total, -1.82, -0.182, 88.998])
    plain = io.BytesIO()
    wb.save(plain)
    locked = io.BytesIO()
    msoffcrypto.format.ooxml.OOXMLFile(io.BytesIO(plain.getvalue())).encrypt(password, locked)
    return locked.getvalue()


def make_zoho() -> bytes:
    """Two days of the Zoho export; each item line repeats the invoice Total."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.append(
        [
            "Invoice Date",
            "Invoice ID",
            "Invoice Number",
            "Invoice Status",
            "Customer Name",
            "Customer Name (Secondary Language)",
            "Company Name",
            "Total",
            "Item Name",
            "Quantity",
            "Item Price",
            "SKU",
            "Sales person",
        ]
    )
    head = ["4681", "INV-001071", "Draft", "CAUSEWAY", "", "CAUSEWAY", 450.493]
    for item, qty in [
        ("Preconfigured- Batelco BD 9.5 Starter Pack", 17),
        ("EVD - KFC 1 - 056194010436028", 230.5),
    ]:
        ws.append([datetime(2026, 10, 18), *head, item, qty, 8.636, "SKU1", "Ali"])
    ws.append(
        [
            datetime(2026, 10, 19),
            "4682",
            "INV-001072",
            "Draft",
            "CAUSEWAY",
            "",
            "CAUSEWAY",
            438.992,
            "eSIM 9.5(Kunooz)",
            12,
            9.5,
            "SKU2",
            "Ali",
        ]
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def run(params=None, afs=None, password=AFS_PASSWORD):
    ctx = RunContext(
        files=[
            InputFile("afs", "afs.xlsx", afs or make_afs(), password),
            InputFile("zoho", "zoho.xlsx", make_zoho()),
        ],
        params={
            "sales_date": "2026-10-18",
            "cash": "100.500",
            "cash_deposit_date": "2026-10-18",
            "pie": "215.000",
            **(params or {}),
        },
        rules={},
    )
    return ctx, CausewaySalesReconciliation().run(ctx)


def test_marketing_fund_matches_the_worked_example():
    ctx, result = run()
    assert [(m.label, m.value, m.tone) for m in result.metrics] == [
        ("Zoho total", "450.493", "neutral"),
        ("Cash + AFS + PIE", "406.500", "neutral"),
        ("Variance", "43.993", "good"),
    ]
    assert result.warnings == []
    assert result.tables[2].rows == [["03221330", "91.000", "-1.820", "-0.182", "88.998"]]
    assert result.tables[3].rows == [["INV-001071", "Draft", "450.493"]]
    assert [o.name for o in ctx.outputs] == [
        "CreditNoteJV-.xlsx",
        "RECEIPTS.xlsx",
        "causeway_reconciliation_20261018.xlsx",
        "zoho.xlsx",
    ]


def test_zoho_file_comes_back_with_formatted_sheet():
    ctx, _ = run()
    zoho = output(ctx, "zoho.xlsx")
    assert zoho.name == "zoho.xlsx"  # same name as the upload
    wb = openpyxl.load_workbook(io.BytesIO(zoho.data))
    assert wb.sheetnames == ["Invoice", "Formatted"]
    assert wb["Invoice"].max_column == 13  # original sheet untouched
    rows = list(wb["Formatted"].iter_rows(values_only=True))
    assert [c.lower() for c in rows[0]] == [c.lower() for c in ZOHO_COLUMNS]
    # Two lines on 18/10, an empty row, then the 19/10 line.
    assert [r[1] for r in rows[1:]] == ["INV-001071", "INV-001071", None, "INV-001072"]
    assert rows[1][4] == 450.493 and rows[1][9] == "Ali"


def test_wrong_afs_date_and_negative_fund_warn():
    _, result = run(params={"cash": "200"}, afs=make_afs(day="17-OCT-2026"))
    assert result.metrics[2].value == "-55.507" and result.metrics[2].tone == "bad"
    assert len(result.warnings) == 2


def test_date_missing_from_zoho():
    with pytest.raises(WorkflowInputError, match="No Zoho invoices for 02/11/2026"):
        run(params={"sales_date": "2026-11-02", "cash_deposit_date": "2026-11-02"})


def test_bad_amount():
    with pytest.raises(WorkflowInputError, match="Cash amount"):
        run(params={"cash": "abc"})


def test_afs_total_check_passes():
    check = CausewaySalesReconciliation().check_file(
        InputFile("afs", "afs.xlsx", make_afs(), AFS_PASSWORD), {}
    )
    assert check.ok
    assert check.message == "Total matches: 10 transactions = 91.000 BHD"


def test_afs_total_not_matching_rows_fails_check_and_run():
    afs = make_afs(total=92)
    check = CausewaySalesReconciliation().check_file(
        InputFile("afs", "afs.xlsx", afs, AFS_PASSWORD), {}
    )
    assert not check.ok
    assert check.message == (
        "Terminal 03221330: the Tran Amt rows add up to 91.000, but the Total row says 92.000."
    )
    with pytest.raises(WorkflowInputError, match="does not add up"):
        run(afs=afs)


def test_next_day_afs_is_filled_and_listed():
    # The last two sales (11.000 each, batch 999) are dated after the sales date, on two days.
    afs = make_afs(later=("19-OCT-2026", "20-OCT-2026"))
    check = CausewaySalesReconciliation().check_file(
        InputFile("afs", "afs.xlsx", afs, AFS_PASSWORD), {"sales_date": "2026-10-18"}
    )
    assert check.ok and check.fills == {"afs_next_day": "22.000"}
    assert check.notice is not None
    assert "batch 999" in check.notice.message
    assert "more than one day: 19/10/2026, 20/10/2026" in check.notice.message
    assert check.notice.table.rows == [
        ["19/10/2026", "999", "03221330", "11.000"],
        ["20/10/2026", "999", "03221330", "11.000"],
    ]


def test_no_date_yet_no_fill():
    check = CausewaySalesReconciliation().check_file(
        InputFile("afs", "afs.xlsx", make_afs(), AFS_PASSWORD), {}
    )
    assert check.fills == {} and check.notice is None


def test_variance_uses_next_and_previous_day_but_receipt_keeps_full_total():
    afs = make_afs(later=("19-OCT-2026",))
    _, result = run(params={"afs_next_day": "11.000", "afs_previous_day": "5.000"}, afs=afs)
    # AFS for the sales date = 91 - 11 + 5 = 85; payments 100.5 + 85 + 215 = 400.5.
    assert result.metrics[1].value == "400.500"
    assert result.metrics[2].value == "49.993"
    assert result.warnings == []
    receipt = next(t for t in result.tables if t.title == "Receipt (AFS)")
    assert any("91.000" in str(row) for row in receipt.rows)


def test_changed_next_day_amount_warns_but_runs():
    afs = make_afs(later=("19-OCT-2026",))
    _, result = run(params={"afs_next_day": "5.000"}, afs=afs)
    assert result.metrics[2].value == "48.993"  # 450.493 - (100.5 + 86 + 215)
    assert len(result.warnings) == 1
    assert "entered as 5.000, but the AFS file has 11.000" in result.warnings[0]


def test_cash_deposit_date_is_shown_on_cash_receipt():
    _, result = run()
    cash = next(t for t in result.tables if t.title == "Receipt (Cash)")
    assert cash.note == "Cash deposit date: 18/10/2026 (same day as the sales date)."
    assert cash.note_tone == "neutral"

    _, result = run(params={"cash_deposit_date": "2026-10-20"})
    cash = next(t for t in result.tables if t.title == "Receipt (Cash)")
    assert cash.note == "Cash deposit date: 20/10/2026, 2 days after the sales date (18/10/2026)."
    assert cash.note_tone == "warn"


def test_cash_deposit_date_before_sales_date_is_refused():
    with pytest.raises(WorkflowInputError, match="same day or later"):
        run(params={"cash_deposit_date": "2026-10-17"})


def test_wrong_password():
    from app.core.excel import SpreadsheetError

    with pytest.raises(SpreadsheetError, match="password"):
        run(password="1111")


def page3_zoho() -> dict:
    """The 03/10/2026 lines from procedure page 3, plus a KFC line and another day."""
    import polars as pl

    rows = [
        ("2026-10-03", "Preconfigured- Batelco BD 9.5 Starter Pack", "STPKVOICEBD9.5", 17, 8.636),
        ("2026-10-03", "Blank Batelco 9.5(kunooz)", "VANILASIM9.5", 6, 9.5),
        ("2026-10-03", "Blank Batelco - 10.5 (EVD2)", "VANILASIM  10.5", 3, 10.5),
        ("2026-10-03", "Blank Batelco 12(EVD 2)", "VANILASIM 12", 4, 12.0),
        ("2026-10-03", "EVD - KFC 2 - 056194020436028", "EVD mPOS KFC 2", 78.5, 1.0),
        ("2026-10-04", "Preconfigured- Batelco BD 9.5 Starter Pack", "STPKVOICEBD9.5", 99, 8.636),
    ]
    schema = ["Invoice Date", "Item Name", "SKU", "Quantity", "Item Price"]
    return {"Invoice": pl.DataFrame(rows, schema=schema, orient="row")}


def test_direct_sales_invoice_matches_page_3():
    from datetime import date

    from app.workflows.finance.causeway_sales.invoice import (
        DEFAULT_INVOICE_RULES,
        DEFAULT_ITEM_TYPES,
        day_lines,
        direct_sales_invoice,
        table_rows,
    )

    lines = day_lines(page3_zoho(), date(2026, 10, 3), DEFAULT_ITEM_TYPES)
    assert [z.kind for z in lines] == ["Normal", "Kunooz", "EVD", "EVD", "KFC"]
    rows = table_rows(direct_sales_invoice(lines, DEFAULT_INVOICE_RULES))
    assert [r[2:7] for r in rows] == [
        ["STPKVOICEBD9.5", "", "17", "8.636", "146.812"],
        ["VANILASIM", "Batelco Vanilla Sim", "13", "0.000", "0.000"],  # EVD + Kunooz, no KFC
        ["0108002", "Commission Income - Batelco Starter Pack", "6", "1.818", "10.909"],
        ["ESIM9.5", "eSIM 9.5(Kunooz)", "6", "From ERP", "From ERP"],
    ]
    assert "VANILASIM9.5" in rows[3][7]  # the Zoho SKU is shown for reference


def test_direct_sales_invoice_without_kunooz():
    from datetime import date

    from app.workflows.finance.causeway_sales.invoice import (
        DEFAULT_INVOICE_RULES,
        DEFAULT_ITEM_TYPES,
        day_lines,
        direct_sales_invoice,
    )

    lines = day_lines(page3_zoho(), date(2026, 10, 4), DEFAULT_ITEM_TYPES)
    invoice = direct_sales_invoice(lines, DEFAULT_INVOICE_RULES)
    assert [(ln.code, ln.qty) for ln in invoice] == [("STPKVOICEBD9.5", 99)]


def test_credit_note_matches_page_4():
    from datetime import date

    from app.workflows.finance.causeway_sales.credit_note import (
        DEFAULT_CREDIT_NOTE_RULES,
        calc_rows,
        credit_note,
        jv_rows,
    )
    from app.workflows.finance.causeway_sales.invoice import DEFAULT_ITEM_TYPES, day_lines

    note = credit_note(day_lines(page3_zoho(), date(2026, 10, 3), DEFAULT_ITEM_TYPES))
    assert [r[4] for r in calc_rows(note)] == ["31.500", "48.000", "78.500", "158.000"]
    assert jv_rows(note, DEFAULT_CREDIT_NOTE_RULES) == [
        [1, "CAUSEWAY", "Causeway Cash Sales", "158.000", "0.000"],
        [2, "056194020436028", "CAUSEWAY CASH SALES 2", "0.000", ""],  # no credit amount
    ]
    assert note.warnings == []


def test_credit_note_terminal_from_item_name_and_price_from_sku():
    from decimal import Decimal as D

    from app.workflows.finance.causeway_sales.credit_note import credit_note
    from app.workflows.finance.causeway_sales.invoice import ZohoLine

    note = credit_note(
        [
            # SKU price wins over the Zoho Item Price; the doubled L doesn't matter.
            ZohoLine("Blank Batelco - 10.5 (EVD1)", "VANILLASIM 10.5", D(2), D(99), "EVD"),
            # KFC 1's SKU has no terminal number, so it comes from the item name.
            ZohoLine(
                "EVD - KFC 1 - 056194010436028", "EVD mPOS - KFC 1 -", D("230.5"), D(1), "KFC"
            ),
        ]
    )
    assert note.total == D("251.500")
    assert note.terminals == ["056194010436028"]


def test_receipts_match_page_5():
    _, result = run()
    tables = {t.title: t.rows for t in result.tables}
    bank = "AUB BHD Account 0016-525530-001"
    afs_name = "ARAB FINANCIAL SERVICES CO. B.S.C.(CLOSED)"
    pie_name = "PAYMENT INTERNATIONAL ENTERPRISE-QR Charges"
    pie_cheque = "MSWIPE 18/10/2026"
    assert tables["Receipt (Cash)"] == [
        ["CAUSEWAY", "Causeway Cash Sales", "100.500", "", "Cash in Hand - Bahraini Dinars"],
    ]
    assert tables["Receipt (AFS)"] == [
        ["CAUSEWAY", "Causeway Cash Sales", "91.000", "AFS 999", bank],
        ["AFS", afs_name, "-2.002", "AFS 999", bank],  # 91 x 2.2%
    ]
    assert tables["Receipt (PIE)"] == [
        ["CAUSEWAY", "Causeway Cash Sales", "215.000", pie_cheque, bank],
        ["PAYME", pie_name, "-3.548", pie_cheque, bank],  # 215 x 1.65%
    ]
    assert "Receipts" not in tables


def test_afs_commission_rate_mismatch_warns():
    ctx = RunContext(
        files=[
            InputFile("afs", "afs.xlsx", make_afs(), AFS_PASSWORD),
            InputFile("zoho", "zoho.xlsx", make_zoho()),
        ],
        params={
            "sales_date": "2026-10-18",
            "cash": "100.500",
            "cash_deposit_date": "2026-10-18",
            "pie": "215.000",
        },
        rules={"receipts": {"afs_commission_percent": "2"}},
    )
    result = CausewaySalesReconciliation().run(ctx)
    assert any("AFS commission" in w for w in result.warnings)


def test_variance_credit_note_puts_all_on_marketing_fund_until_split():
    _, result = run()
    note = next(t for t in result.tables if t.title == "Credit Note (Variance)")
    assert note.rows == [
        [1, "CAUSEWAY", "Causeway Cash Sales", "0.000", "43.993"],
        [2, "201002004001", "Batelco marketing fund", "43.993", "0.000"],
        [3, "501002056", "Business Promotion", "0.000", "0.000"],
        ["", "", "Total", "43.993", "43.993"],
    ]
    assert note.note.startswith("Narration: Causeway Cash Sales 18/10/2026")
    assert result.split.total == "43.993" and not result.split.confirmed
    assert [p.code for p in result.split.parts] == ["201002004001", "501002056"]


def test_no_variance_credit_note_when_variance_is_negative():
    _, result = run(params={"cash": "200"}, afs=make_afs(day="17-OCT-2026"))
    assert result.split is None
    assert all(t.title != "Credit Note (Variance)" for t in result.tables)


def test_variance_credit_note_is_the_tab_after_receipt_pie():
    _, result = run()
    titles = [t.title for t in result.tables]
    assert titles[-2:] == ["Receipt (PIE)", "Credit Note (Variance)"]


def output(ctx, name: str):
    return next(o for o in ctx.outputs if o.name == name)


def erp_sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_credit_note_jv_file_matches_the_pact_template():
    ctx, result = run()
    ws = erp_sheet(output(ctx, "CreditNoteJV-.xlsx").data)
    assert [c.value for c in ws[1]] == [
        "Doc Date",
        "DocPrefix",
        "Doc No",
        "Account Code",
        "Account Name",
        "Debit Amount",
        "Credit Amount",
        "Terminal ID",
        "Narration",
        "Remarks",
        "Business Unit",
    ]
    text = "Causeway Cash Sales Marketing fund against sales date 18/10/2026"
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert rows == [
        [
            datetime(2026, 10, 18),
            "YQ-26-10",
            None,
            "160700-1",
            "CAUSEWAY",
            None,
            43.993,
            None,
            text,
            text,
            "YQ",
        ],
        [
            datetime(2026, 10, 18),
            "YQ-26-10",
            None,
            "201002004001",
            "Batelco marketing fund",
            43.993,
            None,
            None,
            text,
            text,
            "YQ",
        ],
    ]  # Business Promotion is 0.000, so it has no line
    assert ws["C2"].fill.fgColor.rgb == "FFFFFF00"  # Doc No left for the accountant
    assert ws.auto_filter.ref == "A1:K3"
    assert next(t for t in result.tables if t.title == "Credit Note (Variance)").file == (
        "CreditNoteJV-.xlsx"
    )


def test_credit_note_jv_file_after_a_split_has_both_debits():
    from decimal import Decimal

    ctx, result = run()
    result_dict = asdict(result)
    outputs = CausewaySalesReconciliation().apply_split(
        result_dict,
        {"marketing_fund": Decimal("40.000"), "business_promotion": Decimal("3.993")},
        ctx.params,
        {},
    )
    ws = erp_sheet(next(o for o in outputs if o.name == "CreditNoteJV-.xlsx").data)
    assert [(r[3].value, r[5].value, r[6].value) for r in ws.iter_rows(min_row=2)] == [
        ("160700-1", None, 43.993),
        ("201002004001", 40, None),
        ("501002056", 3.993, None),
    ]


def test_receipts_file_for_the_cash_deposit():
    ctx, result = run(params={"cash_deposit_date": "2026-10-20"})
    wb = openpyxl.load_workbook(io.BytesIO(output(ctx, "RECEIPTS.xlsx").data))
    assert wb.sheetnames == ["RECEIPTS"]
    ws = wb["RECEIPTS"]
    assert [c.value for c in ws[1]] == [
        "Doc Date",
        "DocPrefix",
        "Doc No",
        "Business Unit",
        "Cash/Bank GL Code",
        "Cash/Bank GL Name",
        "Account Code",
        "Account Name",
        "Property",
        "Unit",
        "Terminal ID",
        "Amount",
        "Reference",
        "Description",
        "CHEQUE DATE",
        "CHEQUE NUMBER",
    ]
    assert [c.value for c in ws[2]] == [
        datetime(2026, 10, 20),
        "YQ-26-10-",
        None,
        "YQ",
        "101001004021",
        "AUB BHD Account 0016-525530-001",
        "101001001001",
        "Cash in Hand - Bahraini Dinars",
        "-",
        "-",
        "-",
        100.5,
        "CASH DEPOSIT",
        "Causeway Cash sales 18/10/2026",
        "=A2",
        '=CONCATENATE(N2," ",TEXT(O2,"DD/MM/YYYY")," ",C2)',
    ]
    assert ws.max_row == 2 and ws.auto_filter.ref == "A1:P2"
    cash_tab = next(t for t in result.tables if t.title == "Receipt (Cash)")
    assert cash_tab.file == "RECEIPTS.xlsx"


def test_no_receipts_file_without_cash():
    ctx, result = run(params={"cash": "0"})
    assert "RECEIPTS.xlsx" not in [o.name for o in ctx.outputs]
    assert next(t for t in result.tables if t.title == "Receipt (Cash)").file == ""
