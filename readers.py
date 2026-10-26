"""Read the AFS settlement report and the Zoho invoice export."""
import io
import re
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import msoffcrypto
import openpyxl
from msoffcrypto.exceptions import InvalidKeyError

MILS = Decimal("0.001")


class ReadError(Exception):
    """A problem the employee can fix (wrong password, wrong file, missing date)."""


def bhd(value):
    """Turn 91, '8.500 BHD' or '-0.170' into a Decimal rounded to 3 places."""
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(MILS, ROUND_HALF_UP)
    text = re.sub(r"[^0-9.\-]", "", str(value))
    if not text:
        return Decimal("0")
    return Decimal(text).quantize(MILS, ROUND_HALF_UP)


def _open_workbook(data, password=None):
    stream = io.BytesIO(data)
    try:
        office = msoffcrypto.OfficeFile(stream)
        encrypted = office.is_encrypted()
    except Exception:
        encrypted = False  # plain .xlsx is a zip, not an OLE container

    if encrypted:
        if not password:
            raise ReadError("This Excel file is password protected. Enter its password.")
        decrypted = io.BytesIO()
        try:
            office.load_key(password=password)
            office.decrypt(decrypted)
        except InvalidKeyError:
            raise ReadError("Wrong password for the AFS Excel file.")
        except Exception:
            raise ReadError("Could not unlock the AFS Excel file. Check the password.")
        stream = decrypted

    stream.seek(0)
    try:
        return openpyxl.load_workbook(stream, data_only=True, read_only=True)
    except Exception:
        raise ReadError("The file could not be opened as an Excel (.xlsx) workbook.")


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d-%b-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def read_afs(data, password):
    """Sum the 'Total' row of every terminal block in the AFS report.

    Layout (sheet MerchantAUBTerminalWise): 'Terminal ID:...' row, a header row
    with 'Tran Amt', 'Commission', 'VAT', 'Net', the transactions, then a row
    with 'Total' and the block totals.
    """
    wb = _open_workbook(data, password)
    terminals = []
    tran_dates = set()
    tx_count = 0

    for ws in wb.worksheets:
        cols = None
        terminal = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() if v is not None else "" for v in row]
            first = next((c for c in cells if c), "")

            if first.lower().startswith("terminal id"):
                terminal = first.split(":", 1)[-1].strip()
                continue
            if "Tran Amt" in cells:
                cols = {name: cells.index(name) for name in ("Tran Date", "Tran Amt", "Commission", "VAT", "Net") if name in cells}
                continue
            if cols is None:
                continue

            if any(c.lower() == "total" for c in cells):
                terminals.append({
                    "terminal": terminal or ws.title,
                    "tran_amt": bhd(row[cols["Tran Amt"]]),
                    "commission": bhd(row[cols["Commission"]]) if "Commission" in cols else Decimal("0"),
                    "vat": bhd(row[cols["VAT"]]) if "VAT" in cols else Decimal("0"),
                    "net": bhd(row[cols["Net"]]) if "Net" in cols else Decimal("0"),
                })
                cols = None
                continue

            if "Tran Date" in cols:
                d = _parse_date(row[cols["Tran Date"]])
                if d:
                    tran_dates.add(d)
                    tx_count += 1

    if not terminals:
        raise ReadError("No 'Total' row was found in the AFS file. Is this the AFS settlement report?")

    total = {k: sum((t[k] for t in terminals), Decimal("0")) for k in ("tran_amt", "commission", "vat", "net")}
    return {
        "terminals": terminals,
        "dates": sorted(tran_dates),
        "transactions": tx_count,
        **total,
    }


def read_zoho_total(data, wanted):
    """Total invoice amount for one date from the Zoho invoice export.

    Each item line repeats its invoice's Total, so each invoice is counted once.
    """
    wb = _open_workbook(data)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(rows, [])]
    for name in ("Invoice Date", "Invoice Number", "Total"):
        if name not in header:
            raise ReadError(f"The Zoho file has no '{name}' column. Is this the Zoho invoice export?")
    i_date, i_number, i_total = header.index("Invoice Date"), header.index("Invoice Number"), header.index("Total")
    i_status = header.index("Invoice Status") if "Invoice Status" in header else None

    invoices = {}
    dates_in_file = set()
    lines = 0
    for row in rows:
        d = _parse_date(row[i_date])
        if not d:
            continue
        dates_in_file.add(d)
        if d != wanted:
            continue
        lines += 1
        number = str(row[i_number] or "").strip()
        if number not in invoices:
            invoices[number] = {
                "number": number,
                "status": row[i_status] if i_status is not None else "",
                "total": bhd(row[i_total]),
            }

    if not invoices:
        if dates_in_file:
            span = f"{min(dates_in_file):%d/%m/%Y} to {max(dates_in_file):%d/%m/%Y}"
            raise ReadError(f"No Zoho invoices for {wanted:%d/%m/%Y}. The file covers {span}.")
        raise ReadError("The Zoho file has no invoice rows.")

    return {
        "invoices": list(invoices.values()),
        "lines": lines,
        "total": sum((inv["total"] for inv in invoices.values()), Decimal("0")),
    }
