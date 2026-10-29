"""Reading and writing spreadsheets. Workflows use these helpers instead of calling the libraries.

Reading:  Polars + fastexcel (fast, handles .xlsx .xlsm .xls .xlsb .ods), Polars for .csv,
          msoffcrypto-tool to open password-protected files first.
Writing:  XlsxWriter for new formatted workbooks (write_workbook),
          openpyxl for changing an existing workbook and keeping what is in it (edit_workbook),
          and for filling a fixed layout such as an ERP upload file (fill_template).
"""

import io
from collections.abc import Iterator
from contextlib import contextmanager
from copy import copy
from pathlib import Path

import msoffcrypto
import openpyxl
import openpyxl.utils
import polars as pl
import xlsxwriter

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"}
READABLE_SUFFIXES = EXCEL_SUFFIXES | {".csv"}
# Images are accepted for the user to look at (e.g. amounts in an email screenshot), never processed.
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SpreadsheetError(ValueError):
    """The file can't be read as a spreadsheet. The message is safe to show the user."""


def suffix_of(name: str) -> str:
    return Path(name).suffix.lower()


def is_image(name: str) -> bool:
    return suffix_of(name) in IMAGE_SUFFIXES


def is_encrypted(data: bytes) -> bool:
    try:
        return msoffcrypto.OfficeFile(io.BytesIO(data)).is_encrypted()
    except Exception:
        return False


def decrypt(data: bytes, password: str) -> bytes:
    office = msoffcrypto.OfficeFile(io.BytesIO(data))
    office.load_key(password=password)
    out = io.BytesIO()
    try:
        office.decrypt(out)
    except Exception as exc:
        raise SpreadsheetError("The password for this file is wrong.") from exc
    return out.getvalue()


def _clean(df: pl.DataFrame) -> pl.DataFrame:
    """Drop rows where every cell is empty (Excel exports often have spacer rows)."""
    if df.width == 0:
        return df
    return df.filter(~pl.all_horizontal(pl.all().is_null()))


def read_sheets(data: bytes, name: str, password: str | None = None) -> dict[str, pl.DataFrame]:
    """Read every sheet into a DataFrame, keyed by sheet name. A CSV is one sheet named after it."""
    suffix = suffix_of(name)
    if suffix not in READABLE_SUFFIXES:
        raise SpreadsheetError(f"{name}: only Excel or CSV files can be read.")
    if suffix == ".csv":
        try:
            return {Path(name).stem: _clean(pl.read_csv(io.BytesIO(data), infer_schema_length=0))}
        except Exception as exc:
            raise SpreadsheetError(f"{name}: the CSV file could not be read.") from exc
    if is_encrypted(data):
        if not password:
            raise SpreadsheetError(f"{name} is password-protected. Enter its password.")
        data = decrypt(data, password)
    try:
        sheets = pl.read_excel(io.BytesIO(data), sheet_id=0, engine="calamine")
    except Exception as exc:
        raise SpreadsheetError(f"{name}: the file could not be read as Excel.") from exc
    return {sheet: _clean(df) for sheet, df in sheets.items()}


def sheet_names(data: bytes, name: str, password: str | None = None) -> list[str]:
    return list(read_sheets(data, name, password).keys())


def write_workbook(sheets: dict[str, pl.DataFrame]) -> bytes:
    """A new .xlsx with one sheet per DataFrame: bold header, frozen first row, fitted columns."""
    out = io.BytesIO()
    with xlsxwriter.Workbook(out, {"in_memory": True}) as wb:
        for sheet, df in sheets.items():
            df.write_excel(
                workbook=wb,
                worksheet=sheet[:31],
                autofit=True,
                header_format={"bold": True, "bg_color": "#E8EEF4", "border": 1},
                freeze_panes=(1, 0),
                table_style=None,
            )
    return out.getvalue()


@contextmanager
def edit_workbook(data: bytes, name: str) -> Iterator[openpyxl.Workbook]:
    """Open an existing .xlsx/.xlsm for changes, keeping its sheets and formatting.

    with edit_workbook(data, name) as wb:
        add_sheet(wb, "Cleaned", df)
    result = wb.saved   # bytes of the changed workbook
    """
    if suffix_of(name) not in {".xlsx", ".xlsm"}:
        raise SpreadsheetError(f"{name}: only .xlsx or .xlsm files can be edited in place.")
    wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=suffix_of(name) == ".xlsm")
    yield wb
    out = io.BytesIO()
    wb.save(out)
    wb.saved = out.getvalue()  # type: ignore[attr-defined]


def fill_template(template: Path, rows: list[list], first_row: int = 2) -> bytes:
    """A new workbook from a template file (header row + one styled example row at first_row):
    the rows are written from first_row down, each cell styled like the example row's cell in the
    same column (font, fill, border, number format), and the header filter is stretched over them.
    None leaves a cell empty. A callable value gets the row number and returns the cell value, for
    formulas that point at their own row (lambda r: f"=A{r}"). Excel recalculates formulas when the
    file is opened. The template itself is never changed."""
    wb = openpyxl.load_workbook(template)
    ws = wb.active
    styles = [
        (copy(c.font), copy(c.fill), copy(c.border), c.number_format, copy(c.alignment))
        for c in ws[first_row]
    ]
    if ws.max_row >= first_row:
        ws.delete_rows(first_row, ws.max_row - first_row + 1)
    for r, values in enumerate(rows, start=first_row):
        for col, (font, fill, border, number_format, alignment) in enumerate(styles, start=1):
            value = values[col - 1] if col <= len(values) else None
            cell = ws.cell(row=r, column=col, value=value(r) if callable(value) else value)
            cell.font, cell.fill, cell.border = copy(font), copy(fill), copy(border)
            cell.number_format, cell.alignment = number_format, copy(alignment)
    if ws.auto_filter.ref:
        last = openpyxl.utils.get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A{first_row - 1}:{last}{max(first_row - 1 + len(rows), first_row)}"
    wb.calculation.fullCalcOnLoad = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def add_sheet(
    wb: openpyxl.Workbook, title: str, df: pl.DataFrame, header_like: str | None = None
) -> None:
    """Append a DataFrame as a new sheet (header in row 1) to a workbook opened with edit_workbook.

    header_like: an existing sheet whose header style (font, fill, border) the new header copies.
    An all-empty row in df is written as an empty row.
    """
    ws = wb.create_sheet(title=title[:31])
    ws.append(df.columns)
    for row in df.iter_rows():
        ws.append(list(row))
    if header_like is not None:
        source = wb[header_like][1]
        for i, cell in enumerate(ws[1]):
            if i < len(source):
                cell.font = copy(source[i].font)
                cell.fill = copy(source[i].fill)
                cell.border = copy(source[i].border)
    for i, name in enumerate(df.columns, start=1):
        longest = max([len(str(name)), *(len(str(v)) for v in df[name] if v is not None)])
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(longest + 2, 60)
    ws.freeze_panes = "A2"
